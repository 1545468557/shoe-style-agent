"""导出：Word 设计方案 + Excel（BOM / 尺寸表 / 打样单）。

口径（写进文件里，导出后也守得住）：
- 界面上与文件里都必须写明**「AI 生成建议 · 数据为仿真实例」**；
- 效果图是 **AI 生成示意图**（不是实拍、不是实物照片）；
- 成本是**估算值**；打样是**仿真流程 + 预留标准对接位**。
"""

from __future__ import annotations

import time
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import get_config
from .gates import status as gate_status
from .library import LibraryError, find
from .models import (
    Brief,
    MaterialPick,
    PatternDesign,
    PatternIteration,
    PatternPick,
    Project,
    Sampling,
    StyleDirection,
)
from .workflows.bom import build_bom, compute_size_spec, pattern_line

DISCLAIMER = (
    "本方案由 AI 辅助生成：款式图为 AI 生成示意图（示意效果，非实物照片、非真人实拍）；"
    "面料、辅料与版型数据来自**AI 生成建议（仿真实例）**；成本为估算值；打样为仿真流程（已预留标准对接位）。"
    "涉及投产前请由设计师、版师与采购复核。"
)


def _size_table(spec: dict) -> tuple[list[dict], list[dict], str]:
    """尺寸表的行与档位（**部位名与单位都由后端按品类给出**，导出不再写死女装五部位）。"""
    rows = list(spec.get("rows") or [])
    if not rows:                                              # 兼容旧结构
        for key in ("bust", "waist", "hip", "skirt_length", "sleeve_length"):
            label = (spec.get("labels") or {}).get(key)
            value = spec.get(key)
            if label and value not in (None, ""):
                rows.append({"key": key, "label": label, "value": value, "unit": spec.get("unit", "cm")})
    tiers = [t for t in (spec.get("tiers") or []) if t]
    return rows, tiers, str(spec.get("unit") or "cm")


def _material_desc(pick: MaterialPick, bom_material: dict) -> str:
    """面料描述：示例库条目走库；**AI 生成的面料不再查库**（否则 Word 导出会直接报错）。"""
    try:
        material = find("materials", pick.material_id)
        price = material["price_yuan_per_m"]
        return (
            f"面料：{material['name']}（{material['fiber']}，{material['weight_gsm']}g/m²，"
            f"{price[0]}~{price[1]} 元/米）— 来源：示例库（仿真实例）"
        )
    except LibraryError:
        price = [float(x) for x in (bom_material.get("price_range") or [])]
        money = (
            f"估算 {price[0]:g}~{price[1]:g} {bom_material.get('unit') or '元/计价单位'}"
            if len(price) == 2 and any(price)
            else "价格区间见「成本估算」"
        )
        return (
            f"面料：{pick.note or bom_material.get('name') or 'AI 生成面料'}"
            f"｜来源：AI 生成建议（非真实物料，{money}）"
        )


def _pattern_payload(session: Session, project_id: int) -> dict:
    row = session.scalar(
        select(PatternDesign).where(PatternDesign.project_id == project_id).order_by(PatternDesign.id.desc())
    )
    return dict((row.payload if row else {}) or {})


def _out_dir() -> Path:
    path = get_config().assets_dir.parent / "exports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _gather(session: Session, project_id: int) -> dict:
    project = session.get(Project, project_id)
    if project is None:
        raise ValueError("找不到项目")
    brief = session.scalar(select(Brief).where(Brief.project_id == project_id))
    directions = session.scalars(
        select(StyleDirection).where(StyleDirection.project_id == project_id).order_by(StyleDirection.seq)
    ).all()
    material = session.scalar(select(MaterialPick).where(MaterialPick.project_id == project_id))
    pattern = session.scalar(select(PatternPick).where(PatternPick.project_id == project_id))
    iterations = session.scalars(
        select(PatternIteration)
        .where(PatternIteration.project_id == project_id)
        .order_by(PatternIteration.iteration_no)
    ).all()
    sampling = session.scalar(select(Sampling).where(Sampling.project_id == project_id))
    return {
        "project": project,
        "brief": brief,
        "directions": directions,
        "material": material,
        "pattern": pattern,
        "iterations": iterations,
        "sampling": sampling,
        "bom": build_bom(session, project_id),
        "size_spec": compute_size_spec(session, project_id),
        "gates": gate_status(session, project_id),
    }


def export_excel(session: Session, project_id: int) -> Path:
    """Excel：BOM / 尺寸表 / 打样单 三个工作表。"""
    data = _gather(session, project_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["类别", "名称", "用量/规格", "单价区间", "小计区间"])
    for item in data["bom"]["items"]:
        price = item["price_range"]
        money = (
            "—（不计入成本）"
            if not any(price)
            else f"{price[0]}~{price[1]}"
        )
        sub = item["subtotal"]
        sub_text = "—" if not any(sub) else f"{sub[0]}~{sub[1]}"
        ws.append([item["kind"], item["name"], item["qty"], money, sub_text])
    low, high = data["bom"]["estimated_cost_yuan"]
    ws.append(["合计", "", "", "", f"{low}~{high} 元（估算值）"])
    ws.append(["辅料来源", data["bom"].get("trims_source", ""), "", "", ""])

    spec = data["size_spec"]
    rows, tiers, unit = _size_table(spec)
    ws2 = wb.create_sheet("尺寸表")
    ws2.append(["基准", spec["system"], "", f"单位：{unit}"])
    ws2.append(["部位", "数值", "", ""])
    for row in rows:
        ws2.append([row["label"], row["value"], row.get("unit", unit), ""])
    if tiers:
        keys = [k for k in tiers[0] if k != "size"]
        ws2.append([])
        ws2.append([f"尺码档（估算，单位 {unit}）", *keys])
        for tier in tiers:
            ws2.append([tier.get("size", ""), *[tier.get(k, "") for k in keys]])
    ws2.append([])
    ws2.append(["迭代次数", spec["iterations"], "", ""])
    ws2.append(["说明", "尺寸由代码按品类基准 + 版型参数 + 迭代改动计算，估算值，投产前请版师复核", "", ""])

    ws3 = wb.create_sheet("打样单")
    ws3.append(["项目", str(data["project"].name)])
    ws3.append(["生成时间", time.strftime("%Y-%m-%d %H:%M:%S")])
    direction = next(
        (d for d in data["directions"] if data["material"] and d.id == data["material"].direction_id), None
    )
    ws3.append(["选定方向", direction.name if direction else ""])
    ws3.append(["面料", data["material"].note if data["material"] else ""])
    ws3.append(["版型", data["pattern"].pattern_id if data["pattern"] else ""])
    ws3.append(["工艺", ", ".join(data["pattern"].crafts) if data["pattern"] else ""])
    ws3.append(["打样状态", data["sampling"].state if data["sampling"] else "draft"])
    ws3.append([])
    ws3.append(["说明", DISCLAIMER])

    path = _out_dir() / f"鞋服方案_{project_id}_{time.strftime('%Y%m%d-%H%M%S')}.xlsx"
    wb.save(path)
    return path


def export_word(session: Session, project_id: int) -> Path:
    """Word：设计方案（企划 → 方向（含图）→ 面料 → 版型与迭代 → BOM → 尺寸表 → 打样）。"""
    data = _gather(session, project_id)
    doc = Document()
    doc.add_heading(f"鞋服设计方案 · {data['project'].name}", level=0)
    doc.add_paragraph(f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}　·　类别：{data['project'].category}")

    doc.add_heading("一、企划与解析", level=1)
    if data["brief"]:
        doc.add_paragraph(f"企划原文：{data['brief'].raw_text}")
        parsed = data["brief"].parsed or {}
        category = parsed.get("category") or "未提供"
        style = "、".join(parsed.get("style_keywords") or []) or "未提供"
        user = parsed.get("target_user") or "未提供"
        price = parsed.get("price_band") or parsed.get("cost_ceiling") or "未提供"
        doc.add_paragraph(f"解析结果：品类 {category}｜风格 {style}｜人群 {user}｜价格带 {price}")
        if parsed.get("missing"):
            doc.add_paragraph(f"企划未提及：{'、'.join(parsed['missing'])}")

    doc.add_heading("二、风格方向（AI 生成示意图）", level=1)
    for direction in data["directions"]:
        doc.add_heading(f"{direction.seq}. {direction.name}", level=2)
        doc.add_paragraph(direction.inspiration or "")
        if direction.image_path and Path(direction.image_path).exists():
            try:
                doc.add_picture(direction.image_path, width=Inches(3.2))
                caption = "图：AI 生成示意图（占位图）" if direction.is_placeholder else "图：AI 生成示意图"
                doc.add_paragraph(caption)
            except Exception:
                doc.add_paragraph("（图片插入失败，可另行查看）")
        if direction.silhouette:
            doc.add_paragraph(f"廓形：{direction.silhouette}")

    doc.add_heading("三、面料与版型", level=1)
    if data["material"]:
        doc.add_paragraph(_material_desc(data["material"], data["bom"].get("material") or {}))
    if data["pattern"]:
        spec_now = data["size_spec"]
        payload = _pattern_payload(session, project_id)
        try:
            pattern = find("patterns", data["pattern"].pattern_id)
            ease = {"腰围": pattern.get("waist_ease_cm"), "臀围": pattern.get("hip_ease_cm")}
            detail = pattern_line(str(pattern["name"]), spec_now, {"ease": ease})
            doc.add_paragraph(f"版型：{pattern['name']}｜{detail}— 来源：示例库（仿真实例）")
        except LibraryError:
            detail = pattern_line(str(payload.get("category") or ""), spec_now, payload)
            doc.add_paragraph(
                f"版型：{data['pattern'].pattern_id}｜{detail}"
                "— 来源：AI 生成建议（非工厂纸样 / 非打版文件）"
            )
        doc.add_paragraph(f"工艺：{'、'.join(data['pattern'].crafts) or '未选'}")
        doc.add_paragraph(f"辅料：{data['bom'].get('trims_source', '')}")

    doc.add_heading("四、版型迭代记录", level=1)
    if data["iterations"]:
        for row in data["iterations"]:
            changes = "、".join(f"{key} {value:+g}" for key, value in (row.changes or {}).items())
            doc.add_paragraph(f"第 {row.iteration_no} 版：{changes}")
    else:
        doc.add_paragraph("本次没有迭代，采用版型默认参数。")

    doc.add_heading("五、成本估算（估算值）", level=1)
    for item in data["bom"]["items"]:
        sub = item["subtotal"]
        money = "—（不计入成本）" if not any(sub) else f"{sub[0]}~{sub[1]} 元"
        doc.add_paragraph(f"{item['kind']}｜{item['name']}｜{item['qty']}｜{money}")
    low, high = data["bom"]["estimated_cost_yuan"]
    doc.add_paragraph(f"合计：{low}~{high} 元（估算值；正式核价请由采购确认）")

    doc.add_heading("六、尺寸表", level=1)
    spec = data["size_spec"]
    rows, tiers, unit = _size_table(spec)
    doc.add_paragraph(f"基准：{spec['system']}｜单位：{unit}（迭代 {spec['iterations']} 次）")
    doc.add_paragraph("基准码部位：" + " / ".join(f"{r['label']} {r['value']}{r.get('unit', unit)}" for r in rows))
    if tiers:
        keys = [k for k in tiers[0] if k != "size"]
        table = doc.add_table(rows=1, cols=len(keys) + 1)
        table.style = "Table Grid"
        head = table.rows[0].cells
        head[0].text = "尺码"
        for i, key in enumerate(keys, start=1):
            head[i].text = key
        for tier in tiers:
            cells = table.add_row().cells
            cells[0].text = str(tier.get("size", ""))
            for i, key in enumerate(keys, start=1):
                cells[i].text = str(tier.get(key, ""))
        doc.add_paragraph(f"表内数值单位：{unit}；尺寸由代码按品类基准 + 版型参数计算，估算值，投产前请版师复核。")

    doc.add_heading("七、打样流程（仿真）", level=1)
    doc.add_paragraph(f"当前状态：{data['sampling'].state if data['sampling'] else 'draft'}")
    doc.add_paragraph("说明：本方案为仿真打样流程，不接入真实工厂/PLM 系统，已在数据层预留标准对接位。")

    doc.add_heading("附：口径与免责声明", level=1)
    doc.add_paragraph(DISCLAIMER)

    path = _out_dir() / f"鞋服设计方案_{project_id}_{time.strftime('%Y%m%d-%H%M%S')}.docx"
    doc.save(path)
    return path
